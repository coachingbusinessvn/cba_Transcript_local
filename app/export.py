"""Xuất kết quả: DOCX (bảng subtitle), XLSX, SRT, TXT."""
from pathlib import Path

from .segment import Sentence


def _fmt_time(seconds: float, srt: bool = False) -> str:
    ms = int(round(seconds * 1000))
    h, rem = divmod(ms, 3_600_000)
    m, rem = divmod(rem, 60_000)
    s, ms = divmod(rem, 1000)
    if srt:
        return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"
    return f"{h:02d}:{m:02d}:{s:02d}"


def speaker_name(idx: int | None, names: dict | None = None) -> str:
    if idx is None:
        return ""
    if names and str(idx) in names and names[str(idx)].strip():
        return names[str(idx)].strip()
    return f"Người nói {idx + 1}"


def to_docx(sentences: list[Sentence], out_path: Path, title: str,
            names: dict | None = None, meta: dict | None = None) -> Path:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(title, level=1)
    if meta:
        p = doc.add_paragraph()
        run = p.add_run(" · ".join(f"{k}: {v}" for k, v in meta.items()))
        run.font.size = Pt(9)

    has_speaker = any(s.speaker is not None for s in sentences)
    headers = ["#", "Thời gian"] + (["Người nói"] if has_speaker else []) + ["Nội dung"]
    table = doc.add_table(rows=1, cols=len(headers))
    # Style có thể thiếu trên một số bản python-docx / locale — không được làm hỏng export.
    try:
        table.style = "Light Grid Accent 1"
    except (KeyError, ValueError):
        pass
    for cell, h in zip(table.rows[0].cells, headers, strict=True):
        cell.text = h
        if cell.paragraphs[0].runs:
            cell.paragraphs[0].runs[0].bold = True

    for i, s in enumerate(sentences, 1):
        row = table.add_row().cells
        row[0].text = str(i)
        row[1].text = f"{_fmt_time(s.start)} – {_fmt_time(s.end)}"
        if has_speaker:
            row[2].text = speaker_name(s.speaker, names)
        row[-1].text = s.text

    doc.save(str(out_path))
    return out_path


def to_xlsx(sentences: list[Sentence], out_path: Path,
            names: dict | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Transcript"
    has_speaker = any(s.speaker is not None for s in sentences)
    headers = ["#", "Bắt đầu", "Kết thúc"] + (["Người nói"] if has_speaker else []) + ["Nội dung"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)

    for i, s in enumerate(sentences, 1):
        row = [i, _fmt_time(s.start), _fmt_time(s.end)]
        if has_speaker:
            row.append(speaker_name(s.speaker, names))
        row.append(s.text)
        ws.append(row)
        # openpyxl coi chuỗi bắt đầu bằng "=" là công thức: tên người nói hay
        # nội dung transcript do client gửi có thể bắt đầu bằng "=" và chạy như
        # formula khi mở bằng Excel (formula injection) → ép cell về kiểu string.
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":
                cell.data_type = "s"

    widths = [6, 11, 11] + ([16] if has_speaker else []) + [90]
    for col, w in zip("ABCDE", widths, strict=False):
        ws.column_dimensions[col].width = w
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        row[-1].alignment = wrap
    ws.freeze_panes = "A2"
    wb.save(str(out_path))
    return out_path


def to_srt(sentences: list[Sentence], out_path: Path,
           names: dict | None = None) -> Path:
    lines = []
    for i, s in enumerate(sentences, 1):
        who = speaker_name(s.speaker, names)
        text = f"{who}: {s.text}" if who else s.text
        lines.append(
            f"{i}\n{_fmt_time(s.start, srt=True)} --> {_fmt_time(s.end, srt=True)}\n{text}\n"
        )
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return out_path


def to_txt(sentences: list[Sentence], out_path: Path,
           names: dict | None = None) -> Path:
    lines = []
    for s in sentences:
        who = speaker_name(s.speaker, names)
        prefix = f"[{_fmt_time(s.start)}] " + (f"{who}: " if who else "")
        lines.append(prefix + s.text)
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return out_path
