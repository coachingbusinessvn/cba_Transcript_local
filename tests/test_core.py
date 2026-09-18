import struct
import subprocess
import tempfile
import wave
from pathlib import Path

import pytest

from app.diarize import SpeakerTurn, merge_minor_speakers
from app.export import to_docx, to_srt, to_txt, to_xlsx
from app.glossary import correct, parse
from app.media import read_wav16k_float32
from app.segment import Sentence, assign_speakers, build_sentences
from app.transcribe import Word


def _words(spec):
    # spec: [(start, end, " text"), ...] — text có khoảng trắng đầu như Whisper trả về
    return [Word(s, e, t) for s, e, t in spec]


def test_build_sentences_breaks_on_punctuation():
    words = _words(
        [(0.0, 0.5, " Xin"), (0.5, 1.0, " chào."), (1.1, 1.6, " Tôi"), (1.6, 2.2, " là Huy.")]
    )
    sents = build_sentences(words)
    assert [s.text for s in sents] == ["Xin chào.", "Tôi là Huy."]
    assert sents[0].start == 0.0 and sents[0].end == 1.0


def test_build_sentences_breaks_on_gap():
    words = _words([(0.0, 0.5, " một"), (5.0, 5.5, " hai")])  # gap 4.5s > 2.0s
    assert len(build_sentences(words)) == 2


def test_build_sentences_short_pause_does_not_break():
    # Lấy hơi giữa câu (~1.5s < 2.0s) không được cắt câu
    words = _words([(0.0, 0.5, " tôi"), (2.0, 2.5, " nghĩ")])
    assert len(build_sentences(words)) == 1


def test_build_sentences_soft_breaks_long_clause_on_comma():
    # Câu dài không có dấu chấm; qua mốc mềm (12s) gặp dấu phẩy → ngắt tại đó
    words = [Word(float(i), float(i) + 1.0, f" w{i}") for i in range(13)]
    words.append(Word(13.0, 14.0, " nữa,"))  # span ~14s > 12s, kết bằng phẩy
    words.append(Word(14.0, 15.0, " rồi."))
    sents = build_sentences(words)
    assert len(sents) == 2
    assert sents[0].text.endswith("nữa,")


def test_build_sentences_early_comma_does_not_break():
    # Dấu phẩy sớm (câu còn ngắn, < 12s) không được ngắt
    words = _words([(0.0, 0.5, " ừ,"), (0.5, 1.0, " đúng."), (1.0, 1.5, " ok")])
    sents = build_sentences(words)
    assert sents[0].text == "ừ, đúng."


def test_assign_speakers_by_overlap():
    sents = [Sentence(0.0, 2.0, "a"), Sentence(2.5, 4.0, "b"), Sentence(9.0, 9.5, "c")]
    turns = [SpeakerTurn(0.0, 2.2, 0), SpeakerTurn(2.3, 5.0, 1)]
    assign_speakers(sents, turns)
    assert [s.speaker for s in sents] == [0, 1, 1]  # câu c không overlap → theo câu trước


def test_glossary_exact_replacement_vietnamese_boundary():
    g = parse("nờ vi đi a => NVIDIA\n# comment\nVinamilk")
    assert g.hotwords == "NVIDIA, Vinamilk"
    assert correct("hãng nờ vi đi a ra chip mới", g) == "hãng NVIDIA ra chip mới"


def test_glossary_fuzzy_snap():
    g = parse("Vinamilk")
    assert correct("cổ phiếu Vinamik tăng", g) == "cổ phiếu Vinamilk tăng"


def test_glossary_no_false_positive():
    g = parse("Vinamilk")
    text = "vina cà phê rất ngon"
    assert correct(text, g) == text


def test_exports_create_files():
    sents = [Sentence(0.0, 2.0, "Xin chào.", 0), Sentence(2.0, 4.0, "Chào bạn.", 1)]
    names = {"0": "Anh Huy"}
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        to_docx(sents, d / "a.docx", "Test", names, {"Model": "x"})
        to_xlsx(sents, d / "a.xlsx", names)
        to_srt(sents, d / "a.srt", names)
        to_txt(sents, d / "a.txt", names)
        for f in ["a.docx", "a.xlsx", "a.srt", "a.txt"]:
            assert (d / f).exists() and (d / f).stat().st_size > 0
        srt = (d / "a.srt").read_text(encoding="utf-8")
        assert "00:00:00,000 --> 00:00:02,000" in srt
        assert "Anh Huy: Xin chào." in srt
        assert "Người nói 2: Chào bạn." in srt


def test_settings_roundtrip(monkeypatch, tmp_path):
    from app import config

    monkeypatch.setattr(config, "data_dir", lambda: tmp_path)
    s = config.save_settings({"model": "phowhisper-medium", "bogus": 1})
    assert s["model"] == "phowhisper-medium"
    assert "bogus" not in s
    assert config.load_settings()["model"] == "phowhisper-medium"


def test_merge_minor_speakers_collapses_noise_ids():
    """Hai speaker chính + nhiều id 'nhiễu' ngắn → còn 2 người sau merge."""
    turns = [
        SpeakerTurn(0.0, 10.0, 0),  # major A
        SpeakerTurn(10.0, 10.3, 5),  # noise
        SpeakerTurn(10.3, 20.0, 1),  # major B
        SpeakerTurn(20.0, 20.2, 99),  # noise
        SpeakerTurn(20.2, 30.0, 0),  # major A
        SpeakerTurn(30.0, 40.0, 1),  # major B
    ]
    merged = merge_minor_speakers(turns, min_share=0.05, max_speakers=8)
    ids = {t.speaker for t in merged}
    assert ids == {0, 1}
    # Thời lượng gần như giữ nguyên (chỉ gộp nhãn)
    assert abs(sum(t.end - t.start for t in merged) - 40.0) < 1e-6


def test_merge_minor_speakers_respects_max_speakers():
    turns = [SpeakerTurn(float(i), float(i) + 1.0, i) for i in range(20)]
    merged = merge_minor_speakers(turns, min_share=0.0, max_speakers=3)
    assert len({t.speaker for t in merged}) <= 3


def test_read_wav16k_float32_chunked(tmp_path):
    import wave

    import numpy as np

    from app.media import read_wav16k_float32

    samples = np.array([0, 16384, -16384, 32767, -32768], dtype=np.int16)
    p = tmp_path / "t.wav"
    with wave.open(str(p), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(16000)
        w.writeframes(samples.tobytes())
    # _chunk_frames=2 ép đọc thành 3 khối — kiểm tra ráp khối đúng biên
    out = read_wav16k_float32(p, _chunk_frames=2)
    assert out.dtype == np.float32
    assert np.allclose(out, samples.astype(np.float32) / 32768.0)


def test_read_wav16k_float32_rejects_wrong_rate(tmp_path):
    import wave

    import numpy as np
    import pytest

    from app.media import read_wav16k_float32

    p = tmp_path / "bad.wav"
    with wave.open(str(p), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(8000)
        w.writeframes(np.zeros(10, dtype=np.int16).tobytes())
    with pytest.raises(ValueError):
        read_wav16k_float32(p)


def test_duration_seconds_reads_wav(tmp_path):
    import wave

    import numpy as np

    from app.media import duration_seconds

    p = tmp_path / "one_sec.wav"
    with wave.open(str(p), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(16000)
        w.writeframes(np.zeros(16000, dtype=np.int16).tobytes())  # đúng 1.0s
    d = duration_seconds(p)
    assert d is not None and abs(d - 1.0) < 0.05
    # File không tồn tại → None, không raise
    assert duration_seconds(tmp_path / "khong_ton_tai.mp3") is None


def test_checkpoint_roundtrip(tmp_path):
    from app import checkpoint
    from app.transcribe import Word

    p = tmp_path / "partial.jsonl"
    checkpoint.append_words(p, [Word(0.0, 1.0, " xin"), Word(1.0, 2.0, " chào.")])
    checkpoint.append_words(p, [Word(2.5, 3.0, " Huy")])
    words = checkpoint.load_words(p)
    assert [(w.start, w.end, w.text) for w in words] == [
        (0.0, 1.0, " xin"), (1.0, 2.0, " chào."), (2.5, 3.0, " Huy"),
    ]


def test_checkpoint_skips_corrupt_last_line(tmp_path):
    from app import checkpoint

    p = tmp_path / "partial.jsonl"
    p.write_text(
        '{"start": 0.0, "end": 1.0, "text": " ok"}\n{"start": 1.0, "en',  # crash giữa dòng
        encoding="utf-8",
    )
    words = checkpoint.load_words(p)
    assert len(words) == 1 and words[0].text == " ok"
    # File không tồn tại → list rỗng, không raise
    assert checkpoint.load_words(tmp_path / "missing.jsonl") == []


def test_transcribe_calls_segment_cb(monkeypatch, tmp_path):
    from app import transcribe as tr

    class _W:
        def __init__(self, s, e, t):
            self.start, self.end, self.word = s, e, t

    class _Seg:
        def __init__(self, s, e, words):
            self.start, self.end, self.text = s, e, ""
            self.words = words

    class _FakeModel:
        def transcribe(self, *a, **k):
            return iter([_Seg(0, 1, [_W(0, 0.5, " xin"), _W(0.5, 1, " chào.")]),
                         _Seg(1, 2, [_W(1, 2, " Huy.")])]), None

    batches = []
    words = tr.transcribe(_FakeModel(), tmp_path / "x.wav", 2.0,
                          segment_cb=lambda ws: batches.append(list(ws)))
    assert len(words) == 3
    assert [len(b) for b in batches] == [2, 1]  # mỗi segment 1 lần gọi


def test_pipeline_writes_partial_checkpoint(monkeypatch, tmp_path):
    from app import config, media, pipeline
    from app import transcribe as tr
    from app.transcribe import Word

    monkeypatch.setattr(media, "to_wav16k", lambda i, o: o)
    monkeypatch.setattr(media, "wav_duration_seconds", lambda p: 2.0)
    monkeypatch.setattr(config, "model_is_downloaded", lambda k: True)
    monkeypatch.setattr(tr, "load_model", lambda k: object())

    def fake_transcribe(model, wav, duration, segment_cb=None, **kw):
        ws = [Word(0.0, 1.0, " xin"), Word(1.0, 2.0, " chào.")]
        if segment_cb:
            segment_cb(ws)
        return ws

    monkeypatch.setattr(tr, "transcribe", fake_transcribe)
    result = pipeline.run(tmp_path / "in.wav", tmp_path, "phowhisper-small",
                          enable_diarization=False, enable_punctuation=False)
    assert (tmp_path / "partial.jsonl").exists()
    assert len(result["sentences"]) == 1


def test_long_audio_policy():
    import pytest
    from fastapi import HTTPException

    from app.server import _long_audio_policy

    assert _long_audio_policy(600.0, True) == (True, None)          # 10 phút: giữ nguyên
    assert _long_audio_policy(None, True) == (True, None)           # không đo được: giữ nguyên
    d, note = _long_audio_policy(3 * 3600.0, True)                  # 3 giờ: tắt diarize + note
    assert d is False and note and "phân biệt người nói" in note
    assert _long_audio_policy(3 * 3600.0, False) == (False, None)   # 3 giờ, diarize đã tắt sẵn
    with pytest.raises(HTTPException) as ei:                        # 5 giờ: từ chối
        _long_audio_policy(5 * 3600.0, False)
    assert ei.value.status_code == 400


def test_get_job_recovers_partial_from_disk(monkeypatch, tmp_path):
    import json

    from app import config, server

    monkeypatch.setattr(config, "data_dir", lambda: tmp_path)
    job_id = "deadbeef0001"
    jd = tmp_path / "jobs" / job_id
    jd.mkdir(parents=True)
    (jd / "status.json").write_text(json.dumps({
        "status": "running", "stage": "transcribe", "percent": 40,
        "message": "…", "filename": "hop.mp3", "model": "phowhisper-small",
    }), encoding="utf-8")
    (jd / "partial.jsonl").write_text(
        '{"start": 0.0, "end": 1.0, "text": " Xin chào."}\n', encoding="utf-8")

    job = server._get_job(job_id)
    assert job["status"] == "error"
    assert job.get("partial") is True
    assert job["result"]["sentences"][0]["text"] == "Xin chào."
    assert job["result"]["model"] == "phowhisper-small"
    assert job["result"]["punctuation"]["applied"] is False
    assert "chưa được phục hồi dấu câu" in (job.get("note") or "")


def test_get_job_dead_without_partial(monkeypatch, tmp_path):
    import json

    from app import config, server

    monkeypatch.setattr(config, "data_dir", lambda: tmp_path)
    job_id = "deadbeef0002"
    jd = tmp_path / "jobs" / job_id
    jd.mkdir(parents=True)
    (jd / "status.json").write_text(json.dumps({
        "status": "running", "stage": "model", "percent": 6,
        "message": "…", "filename": "hop.mp3", "model": "phowhisper-small",
    }), encoding="utf-8")
    job = server._get_job(job_id)
    assert job["status"] == "error" and "result" not in job


def test_apply_speaker_overrides():
    from app.segment import Sentence
    from app.server import _apply_speaker_overrides

    sents = [Sentence(0, 1, "a", 0), Sentence(1, 2, "b", 1), Sentence(2, 3, "c", 1)]
    _apply_speaker_overrides(sents, {
        "2": 0,        # hợp lệ → đổi
        "0": 5,        # speaker 5 không tồn tại → bỏ qua
        "99": 1,       # index ngoài biên → bỏ qua
        "bad": 1,      # index rác → bỏ qua
        "1": "x",      # speaker rác → bỏ qua
    })
    assert [s.speaker for s in sents] == [0, 1, 0]
    # None và rỗng không đổ vỡ
    _apply_speaker_overrides(sents, None)
    _apply_speaker_overrides(sents, {})


def test_export_srt_with_speaker_override(monkeypatch, tmp_path):
    from fastapi.testclient import TestClient

    from app import config, server

    monkeypatch.setattr(config, "data_dir", lambda: tmp_path)
    job_id = "abc123def456"  # job_id thật luôn là uuid hex 12 ký tự
    (tmp_path / "jobs" / job_id).mkdir(parents=True)
    server.JOBS[job_id] = {
        "status": "done", "filename": "hop.mp3",
        "result": {"sentences": [
            {"start": 0.0, "end": 1.0, "text": "Xin chào.", "speaker": 0},
            {"start": 1.0, "end": 2.0, "text": "Chào bạn.", "speaker": 1},
        ], "duration": 2.0, "num_speakers": 2, "extracted_audio": None, "model": "x"},
    }
    try:
        # base_url loopback: middleware security_guard từ chối Host không phải
        # 127.0.0.1/localhost (mặc định của TestClient là "testserver" → 403).
        c = TestClient(server.app, base_url="http://127.0.0.1")
        r = c.post(f"/api/jobs/{job_id}/export",
                   json={"format": "srt", "speakers": {"1": 0}})
        assert r.status_code == 200
        srt = (tmp_path / "jobs" / job_id / "hop.srt").read_text(encoding="utf-8")
        assert "Người nói 1: Chào bạn." in srt  # dòng 2 đã đổi từ speaker 1 → 0
    finally:
        server.JOBS.pop(job_id, None)


@pytest.mark.timeout(10)
def test_read_wav16k_float32_handles_truncated_file(tmp_path):
    # Mô phỏng ffmpeg bị kill giữa chừng: header khai 10 giây, dữ liệu thực chỉ 1 giây.
    p = tmp_path / "trunc.wav"
    with wave.open(str(p), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(16000)
        w.writeframes(b"\x00\x01" * 16000)
    raw = bytearray(p.read_bytes())
    i = raw.find(b"data")
    struct.pack_into("<I", raw, i + 4, 16000 * 2 * 10)
    p.write_bytes(bytes(raw))

    samples = read_wav16k_float32(p)

    assert len(samples) == 16000  # trả về đúng phần đọc được, không treo


def test_run_ffmpeg_raises_on_timeout(monkeypatch):
    from app import media

    def fake_run(cmd, **kwargs):
        assert kwargs.get("timeout") == media._FFMPEG_TIMEOUT
        raise subprocess.TimeoutExpired(cmd, kwargs["timeout"])

    monkeypatch.setattr(media.subprocess, "run", fake_run)
    monkeypatch.setattr(media, "ffmpeg_path", lambda: "ffmpeg")

    with pytest.raises(RuntimeError, match="quá"):
        media._run_ffmpeg(["-i", "x.wav", "y.wav"])


def test_finished_job_is_evicted_from_memory(monkeypatch, tmp_path):
    from app import config, pipeline, server

    monkeypatch.setattr(config, "data_dir", lambda: tmp_path)
    monkeypatch.setattr(pipeline, "run", lambda *a, **k: {
        "sentences": [{"start": 0.0, "end": 1.0, "text": "xin chào", "speaker": None}],
        "duration": 1.0,
        "num_speakers": 0,
        "extracted_audio": None,
        "model": "phowhisper-small",
        "punctuation": {"requested": False, "applied": False, "warning": None},
    })
    server.JOBS.clear()
    server.JOBS["aa0000000001"] = {
        "status": "running", "filename": "a.wav", "model": "phowhisper-small",
    }

    server._run_job("aa0000000001", tmp_path / "a.wav", {
        "model": "phowhisper-small", "language": "vi", "diarize": False,
        "punctuate": False, "num_speakers": 0, "glossary": "",
    })

    assert "aa0000000001" not in server.JOBS          # không giữ transcript trong RAM
    assert server._get_job("aa0000000001")["status"] == "done"   # vẫn đọc được từ đĩa


def test_failed_job_is_also_evicted_from_memory(monkeypatch, tmp_path):
    from app import config, pipeline, server

    monkeypatch.setattr(config, "data_dir", lambda: tmp_path)

    def boom(*a, **k):
        raise RuntimeError("ffmpeg lỗi")

    monkeypatch.setattr(pipeline, "run", boom)
    server.JOBS.clear()
    server.JOBS["aa0000000002"] = {
        "status": "running", "filename": "b.wav", "model": "phowhisper-small",
    }

    server._run_job("aa0000000002", tmp_path / "b.wav", {
        "model": "phowhisper-small", "language": "vi", "diarize": False,
        "punctuate": False, "num_speakers": 0, "glossary": "",
    })

    # Quan trọng: job lỗi cũng phải rời RAM, nếu không nó chặn mọi job sau bằng 409
    assert "aa0000000002" not in server.JOBS
    assert server._get_job("aa0000000002")["status"] == "error"


def test_create_job_releases_reservation_when_thread_start_fails(monkeypatch, tmp_path):
    # Nếu Thread.start() ném RuntimeError (vd không tạo được thread mới), chỗ đã
    # đặt phải được nhả — nếu không job "running" kẹt vĩnh viễn, mọi job sau 409.
    from fastapi.testclient import TestClient

    from app import config, media, server

    class ExplodingThread:
        def __init__(self, target, args, daemon):
            pass

        def start(self):
            raise RuntimeError("can't start new thread")

    monkeypatch.setattr(config, "data_dir", lambda: tmp_path)
    monkeypatch.setattr(media, "duration_seconds", lambda path: 1.0)
    monkeypatch.setattr(server.threading, "Thread", ExplodingThread)
    server.JOBS.clear()
    try:
        c = TestClient(server.app, base_url="http://127.0.0.1")
        with pytest.raises(RuntimeError, match="can't start new thread"):
            c.post(
                "/api/jobs",
                files={"file": ("hop.wav", b"RIFF", "audio/wav")},
                data={"model": "phowhisper-small"},
            )
        assert server.JOBS == {}
    finally:
        server.JOBS.clear()


def test_finished_job_stays_in_ram_when_disk_write_fails(monkeypatch, tmp_path):
    # Nếu ghi status.json thất bại (ví dụ hết đĩa), job "done" không được biến mất
    # khỏi RAM — nếu không, kết quả không nằm ở đâu cả (đĩa lẫn RAM).
    from app import config, pipeline, server

    monkeypatch.setattr(config, "data_dir", lambda: tmp_path)
    monkeypatch.setattr(pipeline, "run", lambda *a, **k: {
        "sentences": [{"start": 0.0, "end": 1.0, "text": "xin chào", "speaker": None}],
        "duration": 1.0,
        "num_speakers": 0,
        "extracted_audio": None,
        "model": "phowhisper-small",
        "punctuation": {"requested": False, "applied": False, "warning": None},
    })
    monkeypatch.setattr(server, "_save_job", lambda job_id, job: False)
    server.JOBS.clear()
    server.JOBS["aa0000000003"] = {
        "status": "running", "filename": "a.wav", "model": "phowhisper-small",
    }

    server._run_job("aa0000000003", tmp_path / "a.wav", {
        "model": "phowhisper-small", "language": "vi", "diarize": False,
        "punctuate": False, "num_speakers": 0, "glossary": "",
    })

    assert "aa0000000003" in server.JOBS
    assert server.JOBS["aa0000000003"]["status"] == "done"       # không kẹt ở "running"
    assert server.JOBS["aa0000000003"]["result"]["sentences"]    # kết quả vẫn đọc được qua RAM


def test_failed_job_stays_in_ram_when_disk_write_fails(monkeypatch, tmp_path):
    # Tương tự nhưng cho nhánh lỗi: pop cũng phải được gate theo kết quả ghi đĩa.
    from app import config, pipeline, server

    monkeypatch.setattr(config, "data_dir", lambda: tmp_path)

    def boom(*a, **k):
        raise RuntimeError("ffmpeg lỗi")

    monkeypatch.setattr(pipeline, "run", boom)
    monkeypatch.setattr(server, "_save_job", lambda job_id, job: False)
    server.JOBS.clear()
    server.JOBS["aa0000000004"] = {
        "status": "running", "filename": "b.wav", "model": "phowhisper-small",
    }

    server._run_job("aa0000000004", tmp_path / "b.wav", {
        "model": "phowhisper-small", "language": "vi", "diarize": False,
        "punctuate": False, "num_speakers": 0, "glossary": "",
    })

    assert "aa0000000004" in server.JOBS
    assert server.JOBS["aa0000000004"]["status"] == "error"      # không kẹt ở "running"


def test_sha256_file_matches_hashlib(tmp_path):
    import hashlib

    from app import config

    p = tmp_path / "x.bin"
    p.write_bytes(b"vitranscript" * 1000)
    assert config.sha256_file(p) == hashlib.sha256(p.read_bytes()).hexdigest()


def test_download_rejects_wrong_hash(tmp_path, monkeypatch):
    import io

    from app import diarize

    class FakeResponse(io.BytesIO):
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    monkeypatch.setattr(
        diarize.urllib.request, "urlopen", lambda url, timeout: FakeResponse(b"noi dung sai")
    )
    dest = tmp_path / "model.onnx"

    with pytest.raises(RuntimeError, match="SHA-256"):
        diarize._download("https://example.invalid/m.onnx", dest, "0" * 64)

    assert not dest.exists()                              # không để lại file hỏng
    assert not dest.with_suffix(".onnx.part").exists()    # dọn cả file tạm


def test_download_accepts_matching_hash(tmp_path, monkeypatch):
    import hashlib
    import io

    from app import diarize

    payload = b"noi dung dung"

    class FakeResponse(io.BytesIO):
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    monkeypatch.setattr(
        diarize.urllib.request, "urlopen", lambda url, timeout: FakeResponse(payload)
    )
    dest = tmp_path / "model.onnx"

    diarize._download("https://example.invalid/m.onnx", dest, hashlib.sha256(payload).hexdigest())

    assert dest.read_bytes() == payload


def test_clean_names_drops_non_string_values():
    from app.server import _clean_names

    assert _clean_names({"0": "Huy", "1": 123, "2": None}) == {"0": "Huy"}
    assert _clean_names(None) == {}
    assert _clean_names("khong phai dict") == {}
    assert _clean_names({0: "Huy"}) == {"0": "Huy"}  # key so -> chuoi


def test_load_model_frees_old_model_before_building_new(monkeypatch):
    import app.transcribe as tr

    seen_during_construct = {}

    class FakeWhisperModel:
        def __init__(self, source, **kwargs):
            # Model cu phai da bi giai phong TRUOC khi model moi cap phat RAM
            seen_during_construct["cache"] = dict(tr._MODEL_CACHE)

    fake_module = type("M", (), {"WhisperModel": FakeWhisperModel})
    monkeypatch.setitem(__import__("sys").modules, "faster_whisper", fake_module)

    tr._MODEL_CACHE.clear()
    tr._MODEL_CACHE["phowhisper-large"] = object()
    try:
        tr.load_model("phowhisper-small")

        assert seen_during_construct["cache"] == {}
        assert list(tr._MODEL_CACHE) == ["phowhisper-small"]
    finally:
        tr._MODEL_CACHE.clear()


def test_non_loopback_host_rejected():
    """DNS rebinding: domain attacker resolve về 127.0.0.1 nhưng giữ Host gốc → 403."""
    from fastapi.testclient import TestClient

    from app import server

    c = TestClient(server.app)  # host mặc định "testserver" — không phải loopback
    assert c.get("/").status_code == 403
    ok = TestClient(server.app, base_url="http://127.0.0.1")
    assert ok.get("/").status_code == 200
    ok_local = TestClient(server.app, base_url="http://localhost:8756")
    assert ok_local.get("/").status_code == 200


def test_security_headers_present():
    from fastapi.testclient import TestClient

    from app import server

    r = TestClient(server.app, base_url="http://127.0.0.1").get("/")
    assert r.headers["X-Content-Type-Options"] == "nosniff"
    assert r.headers["X-Frame-Options"] == "DENY"
    assert "frame-ancestors 'none'" in r.headers["Content-Security-Policy"]


def test_cross_site_write_blocked():
    """CSRF: POST mang Origin/Sec-Fetch-Site của site khác bị chặn trước khi tới route."""
    from fastapi.testclient import TestClient

    from app import server

    c = TestClient(server.app, base_url="http://127.0.0.1")
    r = c.post("/api/jobs", headers={"origin": "https://evil.example"})
    assert r.status_code == 403
    r = c.post("/api/jobs", headers={"sec-fetch-site": "cross-site"})
    assert r.status_code == 403
    # Same-origin đi qua guard → request tới route và lỗi validate như bình thường
    r = c.post("/api/jobs", headers={"origin": "http://127.0.0.1:8756"})
    assert r.status_code == 422  # thiếu file/model — nghĩa là đã qua được guard


def test_job_id_path_traversal_rejected(monkeypatch, tmp_path):
    """job_id không đúng uuid-hex-12 (vd '..') bị 404 — không chạm đường dẫn ngoài jobs/."""
    import json

    from fastapi.testclient import TestClient

    from app import config, server

    monkeypatch.setattr(config, "data_dir", lambda: tmp_path)
    # status.json ngoài jobs/ — nếu traversal còn tồn tại, request sẽ đọc file này
    (tmp_path / "status.json").write_text(json.dumps({"status": "done"}), encoding="utf-8")

    c = TestClient(server.app, base_url="http://127.0.0.1")
    assert c.get("/api/jobs/%2E%2E").status_code == 404
    assert c.get("/api/jobs/not-a-job-id").status_code == 404
    assert c.post("/api/jobs/%2E%2E/export", json={"format": "srt"}).status_code == 404


def test_xlsx_stores_equals_text_as_string(tmp_path):
    """Text/tên bắt đầu '=' không được ghi thành formula (Excel formula injection)."""
    from openpyxl import load_workbook

    from app.export import to_xlsx
    from app.segment import Sentence

    evil = "=cmd|'/c calc'!A0"
    sents = [Sentence(0.0, 1.0, evil, 0)]
    out = to_xlsx(sents, tmp_path / "a.xlsx", {"0": evil})

    wb = load_workbook(out)
    text_cells = [c for row in wb.active.iter_rows(min_row=2) for c in row if c.value == evil]
    assert len(text_cells) == 2                     # cột người nói + cột nội dung
    assert all(c.data_type == "s" for c in text_cells)  # string, không phải formula


def test_upload_size_cap_and_num_speakers_clamp(monkeypatch, tmp_path):
    from fastapi.testclient import TestClient

    from app import config, media, pipeline, server

    monkeypatch.setattr(config, "data_dir", lambda: tmp_path)
    monkeypatch.setattr(media, "duration_seconds", lambda p: 1.0)
    monkeypatch.setattr(server, "_MAX_UPLOAD_BYTES", 4)

    class SyncThread:
        def __init__(self, target, args, daemon):
            self.target, self.args = target, args

        def start(self):
            self.target(*self.args)

    monkeypatch.setattr(server.threading, "Thread", SyncThread)
    captured = {}
    monkeypatch.setattr(pipeline, "run", lambda *a, **k: captured.update(k) or {
        "sentences": [{"start": 0.0, "end": 1.0, "text": "xin chào", "speaker": None}],
        "duration": 1.0, "num_speakers": 0, "extracted_audio": None, "model": "x",
        "punctuation": {"requested": False, "applied": False, "warning": None},
    })
    server.JOBS.clear()
    try:
        c = TestClient(server.app, base_url="http://127.0.0.1")
        # Upload vượt cap → 413 + file ghi dở bị xóa
        r = c.post("/api/jobs",
                   files={"file": ("hop.wav", b"RIFFx", "audio/wav")},
                   data={"model": "phowhisper-small"})
        assert r.status_code == 413
        assert not list(tmp_path.rglob("input.wav"))
        # num_speakers khổng lồ bị clamp về _MAX_SPEAKERS trước khi vào pipeline
        monkeypatch.setattr(server, "_MAX_UPLOAD_BYTES", 1 << 30)
        r = c.post("/api/jobs",
                   files={"file": ("hop.wav", b"RIFF", "audio/wav")},
                   data={"model": "phowhisper-small", "num_speakers": "9999"})
        assert r.status_code == 200
        assert captured["num_speakers"] == server._MAX_SPEAKERS
    finally:
        server.JOBS.clear()


def test_download_audio_rejects_path_outside_job_dir(monkeypatch, tmp_path):
    """extracted_audio trỏ ra ngoài job dir (status.json bị sửa) → 404, không đọc file lạ."""
    import json

    from fastapi.testclient import TestClient

    from app import config, server

    monkeypatch.setattr(config, "data_dir", lambda: tmp_path)
    job_id = "aabbccddee01"
    jd = tmp_path / "jobs" / job_id
    jd.mkdir(parents=True)
    outside = tmp_path / "secret.m4a"
    outside.write_bytes(b"secret")
    (jd / "status.json").write_text(json.dumps({
        "status": "done", "filename": "x.mp4", "model": "x",
        "result": {"sentences": [], "extracted_audio": str(outside)},
    }), encoding="utf-8")

    c = TestClient(server.app, base_url="http://127.0.0.1")
    assert c.get(f"/api/jobs/{job_id}/audio").status_code == 404

    # File hợp lệ nằm trong job dir vẫn tải được bình thường
    inside = jd / "input.m4a"
    inside.write_bytes(b"ok")
    (jd / "status.json").write_text(json.dumps({
        "status": "done", "filename": "x.mp4", "model": "x",
        "result": {"sentences": [], "extracted_audio": str(inside)},
    }), encoding="utf-8")
    r = c.get(f"/api/jobs/{job_id}/audio")
    assert r.status_code == 200 and r.content == b"ok"


def test_export_tolerates_extra_keys_in_sentences(monkeypatch, tmp_path):
    """status.json có key lạ trong sentence không được làm export nổ 500."""
    import json

    from fastapi.testclient import TestClient

    from app import config, server

    monkeypatch.setattr(config, "data_dir", lambda: tmp_path)
    job_id = "aabbccddee02"
    jd = tmp_path / "jobs" / job_id
    jd.mkdir(parents=True)
    (jd / "status.json").write_text(json.dumps({
        "status": "done", "filename": "x.mp3", "model": "x",
        "result": {"sentences": [
            {"start": 0.0, "end": 1.0, "text": "a", "speaker": 0, "bogus": 1},
        ], "duration": 1.0, "num_speakers": 1, "extracted_audio": None, "model": "x"},
    }), encoding="utf-8")

    r = TestClient(server.app, base_url="http://127.0.0.1").post(
        f"/api/jobs/{job_id}/export", json={"format": "txt"})
    assert r.status_code == 200
